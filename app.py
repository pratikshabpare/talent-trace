
"""
TalentTrace - Resume & LinkedIn Consistency Checker

Application flow:
1. Receive a resume and LinkedIn URL from the form.
2. Extract text from PDF, DOCX, or TXT.
3. Send the extracted text to n8n for AI parsing.
4. Receive candidate details from n8n/Gemini.
5. Detect the actual columns in the MySQL candidates_new table.
6. Save supported candidate fields without storing resume text
   or the complete AI response.
7. Display the result and delete the temporary uploaded file.
"""

import os
import json
import tempfile

import requests
import mysql.connector

from flask import Flask, render_template, request, flash
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document


# ============================================================
# 1. FLASK CONFIGURATION
# ============================================================

app = Flask(__name__)

app.secret_key = os.getenv(
    "FLASK_SECRET_KEY",
    "local-dev-change-me"
)

UPLOAD_FOLDER = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "uploads"
)

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}

# For testing, start "Listen for test event" in the n8n Webhook node first.
# For production, activate the workflow and use /webhook/resume-profile-check.
N8N_WEBHOOK_URL = os.getenv(
    "N8N_WEBHOOK_URL",
    "http://localhost:5678/webhook-test/resume-profile-check"
).strip()

# The database name and table name are different.
MYSQL_DATABASE = os.getenv(
    "MYSQL_DATABASE",
    "resume_assessment"
)

MYSQL_TABLE = "candidates_new"


# ============================================================
# 2. MYSQL CONNECTION
# ============================================================


def get_db_connection():
    """Connect to the local MySQL database."""

    return mysql.connector.connect(
        host=os.getenv("MYSQL_HOST", "localhost"),
        port=int(os.getenv("MYSQL_PORT", "3306")),
        user=os.getenv("MYSQL_USER", "root"),
        password=os.getenv("MYSQL_PASSWORD", "pratik"),
        database=MYSQL_DATABASE,
        connection_timeout=10
    )



# ============================================================
# 3. RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(file_path, extension):
    """Extract readable text from PDF, DOCX, or TXT."""

    if extension == "pdf":
        reader = PdfReader(file_path)

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        ).strip()

    if extension == "docx":
        document = Document(file_path)

        return "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        ).strip()

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as file:
        return file.read().strip()


# ============================================================
# 4. HELPERS FOR MATCHING EXISTING MYSQL COLUMNS
# ============================================================

def serialize_json(value):
    """Convert lists or dictionaries into JSON strings for MySQL."""

    return json.dumps(
        value if value is not None else [],
        ensure_ascii=False
    )


def normalize_bool(value):
    """Safely normalize booleans returned by n8n/Gemini.

    Strings such as "false", "no", and "0" must not become True.
    Unknown or missing values return None.
    """
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"true", "yes", "1"}:
            return True
        if normalized in {"false", "no", "0"}:
            return False
    return None


def find_existing_column(actual_columns, possible_names):
    """
    Return the actual database column matching one of the
    accepted names, ignoring capitalization.

    Example:
        Actual column: skills
        Accepted names: skills_json, skills
        Result: skills
    """

    lookup = {
        column.lower(): column
        for column in actual_columns
    }

    for name in possible_names:
        if name.lower() in lookup:
            return lookup[name.lower()]

    return None


# ============================================================
# 5. SAVE CANDIDATE TO MYSQL
# ============================================================

def save_candidate(candidate, linkedin_check, filename):
    """
    Save candidate details using the columns that actually
    exist in candidates_new.

    This function does not save:
        - The full extracted resume text
        - The complete AI response

    If a required database column is missing a value, the
    function reports the problem rather than silently changing
    or deleting database data.
    """

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        # ----------------------------------------------------
        # STEP 1: Discover the table's actual column definitions
        # ----------------------------------------------------

        cursor.execute(
            f"SHOW COLUMNS FROM `{MYSQL_TABLE}`"
        )

        column_info = cursor.fetchall()

        if not column_info:
            raise RuntimeError(
                f"The table '{MYSQL_TABLE}' has no columns "
                "or does not exist."
            )

        # Keep the actual column names and their definitions.
        actual_columns = [
            column["Field"]
            for column in column_info
        ]

        # ----------------------------------------------------
        # STEP 2: Prepare candidate values
        # ----------------------------------------------------

        skills_json = serialize_json(
            candidate.get("skills", [])
        )

        education_json = serialize_json(
            candidate.get("education", [])
        )

        experience_json = serialize_json(
            candidate.get("experience", [])
        )

        format_valid = linkedin_check.get("format_valid")
        found_in_resume = linkedin_check.get(
            "url_appears_in_resume"
        )

        # Use None for missing values, which becomes SQL NULL.
        candidate_values = {
            "full_name": candidate.get("full_name"),
            "name": candidate.get("full_name"),

            "email": candidate.get("email"),
            "phone": candidate.get("phone"),

            "skills_json": skills_json,
            "skills": skills_json,

            "education_json": education_json,
            "education": education_json,

            "experience_json": experience_json,
            "experience": experience_json,

            "linkedin_url": linkedin_check.get("normalized_url"),
            "linkedin_profile_url": linkedin_check.get(
                "normalized_url"
            ),

            "linkedin_format_valid": (
                normalize_bool(format_valid)
            ),

            "linkedin_found_in_resume": (
                normalize_bool(found_in_resume)
            ),

            "url_appears_in_resume": (
                normalize_bool(found_in_resume)
            ),

            "resume_filename": filename,
            "filename": filename
        }

        # ----------------------------------------------------
        # STEP 3: Match values to the existing columns
        # ----------------------------------------------------

        insert_values = {}

        for possible_name, value in candidate_values.items():
            actual_name = find_existing_column(
                actual_columns,
                [possible_name]
            )

            if actual_name is not None:
                insert_values[actual_name] = value

        if not insert_values:
            raise RuntimeError(
                f"No supported candidate columns were found "
                f"in '{MYSQL_TABLE}'. Actual columns: "
                f"{actual_columns}"
            )

        # ----------------------------------------------------
        # STEP 4: Detect required columns we cannot populate
        # ----------------------------------------------------

        for column in column_info:
            column_name = column["Field"]
            is_nullable = column["Null"] == "YES"
            has_default = column["Default"] is not None
            is_auto_increment = (
                "auto_increment" in (column["Extra"] or "").lower()
            )

            # Resume text and full parsed JSON are intentionally
            # excluded from storage.
            intentionally_omitted = column_name.lower() in {
                "resume_text",
                "parsed_result_json"
            }

            if (
                not is_nullable
                and not has_default
                and not is_auto_increment
                and column_name not in insert_values
            ):
                if intentionally_omitted:
                    raise RuntimeError(
                        f"Column '{column_name}' is mandatory, "
                        "but this application intentionally does "
                        "not store full resume text or full parsed "
                        "JSON. Make this column nullable in MySQL "
                        "if you want to omit it."
                    )

                raise RuntimeError(
                    f"Column '{column_name}' is required by "
                    "MySQL but is not populated by the application. "
                    "Check the table schema."
                )

        # ----------------------------------------------------
        # STEP 5: Build a parameterized INSERT query
        # ----------------------------------------------------

        column_names = ", ".join(
            f"`{name}`"
            for name in insert_values
        )

        placeholders = ", ".join(
            ["%s"] * len(insert_values)
        )

        insert_sql = (
            f"INSERT INTO `{MYSQL_TABLE}` "
            f"({column_names}) "
            f"VALUES ({placeholders})"
        )

        values = tuple(insert_values.values())

        # Parameterized SQL protects against SQL injection.
        cursor.execute(insert_sql, values)

        # Permanently save the new record.
        connection.commit()

        # Return the generated primary key, if available.
        candidate_id = cursor.lastrowid

        app.logger.info(
            "Candidate saved successfully. Table=%s, ID=%s",
            MYSQL_TABLE,
            candidate_id
        )

        return candidate_id

    except Exception:
        if connection and connection.is_connected():
            connection.rollback()

        # Print the actual reason in the Flask terminal.
        app.logger.exception(
            "Could not save candidate to MySQL table '%s'",
            MYSQL_TABLE
        )

        raise

    finally:
        if cursor:
            cursor.close()

        if connection and connection.is_connected():
            connection.close()


# ============================================================
# 6. HOME PAGE
# ============================================================

@app.route("/")
def index():
    """Display the resume upload form."""

    return render_template(
        "index.html",
        result=None
    )


# ============================================================
# 7. ANALYZE A RESUME
# ============================================================

@app.route("/analyze", methods=["POST"])
def analyze():
    """Process the uploaded resume and display the analysis."""

    uploaded_file = request.files.get("resume")

    submitted_linkedin_url = (
        request.form.get("linkedin_url") or ""
    ).strip()

    # --------------------------------------------------------
    # STEP 1: Validate the uploaded file
    # --------------------------------------------------------

    if not uploaded_file or not uploaded_file.filename:
        flash("Please select a resume file.", "error")

        return render_template(
            "index.html",
            result=None
        ), 400

    filename = secure_filename(uploaded_file.filename)

    if not filename or "." not in filename:
        flash(
            "Please upload a file with a valid extension.",
            "error"
        )

        return render_template(
            "index.html",
            result=None
        ), 400

    extension = filename.rsplit(".", 1)[-1].lower()

    if extension not in ALLOWED_EXTENSIONS:
        flash(
            "Upload a PDF, DOCX, or TXT file.",
            "error"
        )

        return render_template(
            "index.html",
            result=None
        ), 400

    temporary_path = None
    result = None

    try:
        # ----------------------------------------------------
        # STEP 2: Store the uploaded file temporarily
        # ----------------------------------------------------

        with tempfile.NamedTemporaryFile(
            mode="wb",
            suffix="." + extension,
            prefix="resume_",
            dir=UPLOAD_FOLDER,
            delete=False
        ) as temporary_file:

            temporary_path = temporary_file.name
            uploaded_file.save(temporary_file)

        # ----------------------------------------------------
        # STEP 3: Extract resume text
        # ----------------------------------------------------

        resume_text = extract_resume_text(
            temporary_path,
            extension
        )

        if len(resume_text) < 20:
            flash(
                "Not enough text could be extracted. "
                "Scanned PDFs may require OCR.",
                "error"
            )

            return render_template(
                "index.html",
                result=None
            ), 400

        # ----------------------------------------------------
        # STEP 4: Send the resume to n8n
        # ----------------------------------------------------

        webhook_response = requests.post(
            N8N_WEBHOOK_URL,
            json={
                "resume_text": resume_text,
                "linkedin_url": submitted_linkedin_url
            },
            timeout=90
        )

        webhook_response.raise_for_status()

        result = webhook_response.json()

        if not isinstance(result, dict):
            raise ValueError(
                "The n8n response must be a JSON object."
            )

        if result.get("success") is False:
            raise ValueError(
                "The n8n workflow reported a parsing failure."
            )

        candidate = result.get("candidate") or {}
        linkedin_check = result.get("linkedin_check") or {}

        if not isinstance(candidate, dict):
            raise ValueError(
                "Candidate data is missing or invalid."
            )

        if not isinstance(linkedin_check, dict):
            raise ValueError(
                "LinkedIn check data is invalid."
            )

        # ----------------------------------------------------
        # STEP 5: Save candidate details to MySQL
        # ----------------------------------------------------

        candidate_id = save_candidate(
            candidate,
            linkedin_check,
            filename
        )

        result["database"] = {
            "saved": True,
            "candidate_id": candidate_id
        }

        app.logger.info(
            "Resume analyzed and candidate saved."
        )

        return render_template(
            "index.html",
            result=result
        )

    except requests.HTTPError as exc:
        response = exc.response
        status_code = response.status_code if response is not None else 502
        response_body = (
            response.text[:1000] if response is not None else str(exc)
        )

        app.logger.error(
            "n8n returned HTTP %s for %s. Response: %s",
            status_code,
            N8N_WEBHOOK_URL,
            response_body
        )

        if status_code == 404:
            message = (
                "n8n returned 404: the webhook is not registered. "
                "For testing, start 'Listen for test event' and use "
                "/webhook-test/resume-profile-check. For production, "
                "activate the workflow and use /webhook/resume-profile-check."
            )
        else:
            message = (
                f"n8n returned HTTP {status_code}. "
                "Check the n8n execution details and Flask terminal."
            )

        flash(message, "error")
        return render_template("index.html", result=None), 502

    except requests.RequestException as exc:
        app.logger.exception(
            "Could not connect to n8n at %s: %s",
            N8N_WEBHOOK_URL,
            exc
        )

        flash(
            "Could not connect to n8n. Confirm that n8n is running "
            "and N8N_WEBHOOK_URL points to the correct host and port.",
            "error"
        )

        return render_template("index.html", result=None), 502

    except mysql.connector.Error:
        app.logger.exception("MySQL operation failed")

        flash(
            "Resume parsing completed, but MySQL saving failed. "
            "Check the Flask terminal for the exact error.",
            "error"
        )

        if result is not None:
            result["database"] = {
                "saved": False,
                "candidate_id": None
            }

        return render_template(
            "index.html",
            result=result
        ), 500

    except (ValueError, RuntimeError):
        app.logger.exception(
            "Invalid workflow response or database schema"
        )

        flash(
            "Processing could not finish. Check the Flask "
            "terminal for the exact reason.",
            "error"
        )

        if result is not None:
            result["database"] = {
                "saved": False,
                "candidate_id": None
            }

        return render_template(
            "index.html",
            result=result
        ), 500

    except Exception:
        app.logger.exception("Unexpected processing error")

        flash(
            "An unexpected error occurred. Check the Flask "
            "terminal for details.",
            "error"
        )

        return render_template(
            "index.html",
            result=result
        ), 500

    finally:
        # Always attempt to delete the temporary uploaded file.
        if temporary_path and os.path.exists(temporary_path):
            try:
                os.remove(temporary_path)
            except OSError:
                app.logger.exception(
                    "Could not delete temporary resume file"
                )


# ============================================================
# 8. START THE APPLICATION
# ============================================================

if __name__ == "__main__":
    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False
    )
